from datetime import date, datetime
from decimal import Decimal

import pytest
from filelock import FileLock
from openpyxl import load_workbook

from app import config
from app.history import xlsx
from app.models import EmailHistory


@pytest.fixture
def path(tmp_path, monkeypatch):
    p = tmp_path / "data" / "history.xlsx"
    monkeypatch.setattr(config, "HISTORY_XLSX_PATH", str(p))
    monkeypatch.setattr(xlsx, "LOCK_TIMEOUT", 0.2)
    return p


_n = iter(range(1, 10_000))


def add(db, n=1, **kw):
    for _ in range(n):
        db.add(EmailHistory(vendor_name="ABC Pharma", status="SENT", utr=f"UTR{next(_n)}",
                            logged_at=datetime(2025, 9, 15, 10, 0, 0), **kw))
    db.commit()


def sheet(path):
    return [[c.value for c in row] for row in load_workbook(path).active.iter_rows()]


def test_every_row_once_and_twice_adds_nothing(db, path):
    add(db, 3)
    assert xlsx.append_new_rows(db) == 3
    rows = sheet(path)
    assert rows[0] == xlsx.COLUMNS and len(rows) == 4
    assert len({r[6] for r in rows[1:]}) == 3  # three distinct UTRs
    assert xlsx.append_new_rows(db) == 0 and len(sheet(path)) == 4
    add(db, 1)
    assert xlsx.append_new_rows(db) == 1 and len(sheet(path)) == 5


def test_cells_are_indian_format_and_empty_stays_empty(db, path):
    add(db, 1, invoice_no="INV-2305", invoice_amount=Decimal("545060"), amount_paid=Decimal("545060"),
        payment_date=date(2025, 9, 12), tds_amount=None, invoice_deduction=Decimal("0"))
    xlsx.append_new_rows(db)
    cell = dict(zip(xlsx.COLUMNS, sheet(path)[1]))
    assert cell["Invoice Amount"] == cell["Amount Paid"] == "5,45,060.00" and cell["Payment Date"] == "12-Sep-25"
    for empty in ("TDS Amount", "Invoice Deduction", "Reason", "Not Sent Reason", "Sent At", "Error", "To"):
        assert cell[empty] is None


def test_formula_like_text_is_stored_as_text(db, path):
    add(db, 1, reason='=HYPERLINK("http://x")')
    xlsx.append_new_rows(db)
    assert sheet(path)[1][11] == '=HYPERLINK("http://x")'
    assert load_workbook(path).active.cell(2, 12).data_type == "s"


def test_lock_held_raises_nothing_and_next_call_catches_up(db, path):
    add(db, 2)
    with FileLock(f"{path}.lock"):
        assert xlsx.append_new_rows(db) == 0
    assert not path.exists()
    assert xlsx.append_new_rows(db) == 2


def test_file_open_in_excel_leaves_file_intact_then_catches_up(db, path, monkeypatch):
    add(db, 1)
    xlsx.append_new_rows(db)
    add(db, 2)
    real = xlsx.os.replace

    def locked(*a):
        raise PermissionError("open in Excel")

    monkeypatch.setattr(xlsx.os, "replace", locked)
    assert xlsx.append_new_rows(db) == 0
    assert len(sheet(path)) == 2  # old file untouched
    assert not path.with_name(path.name + ".tmp").exists()
    monkeypatch.setattr(xlsx.os, "replace", real)
    assert xlsx.append_new_rows(db) == 2 and len(sheet(path)) == 4


def test_rebuild_restores_deleted_and_corrupted_file(db, path):
    add(db, 3)
    xlsx.append_new_rows(db)
    original = sheet(path)
    path.unlink()
    assert xlsx.rebuild(db) == 3 and sheet(path) == original
    path.write_bytes(b"not an excel file")
    assert xlsx.append_new_rows(db) == 0  # logged, not raised
    assert xlsx.rebuild(db) == 3 and sheet(path) == original
    assert xlsx.append_new_rows(db) == 0  # the mark came back with the file


def test_deleted_file_refills_on_next_append(db, path):
    add(db, 2)
    xlsx.append_new_rows(db)
    path.unlink()
    assert xlsx.append_new_rows(db) == 2


def test_rebuild_endpoint(client, db, path):
    add(db, 2)
    r = client.post("/api/email-history/rebuild-xlsx")
    assert r.status_code == 200 and r.json() == {"rows": 2} and "x-stub" not in r.headers
    with FileLock(f"{path}.lock"):
        assert client.post("/api/email-history/rebuild-xlsx").status_code == 503
