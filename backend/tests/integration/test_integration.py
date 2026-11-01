"""Both tracks together, through the real API: upload -> reconcile -> queue -> send -> history, re-runs,
history completeness, manual review and fixes. SMTP is the fake server from test_sending."""
import json
import smtplib
from collections import Counter
from datetime import date

import pytest
from filelock import FileLock
from openpyxl import load_workbook
from sqlalchemy import func, select

from app import config
from app.history import xlsx
from app.models import AuditLog, EmailHistory, EmailItem, EmailQueue
from app.sending import run_due
from tests.fixtures import special_cases as sc
from tests.fixtures import three_row_example as fx
from tests.test_email_builder import table
from tests.test_sending import smtp  # noqa: F401  (fixture)
from tests.util import files_from, xlsx as make_xlsx

SC_SEND_ON = "2025-09-18"  # special cases are all paid Tue 16 Sep 2025 -> T+2 = Thu 18 Sep


@pytest.fixture
def hist(tmp_path, monkeypatch):
    p = tmp_path / "data" / "email_history.xlsx"
    monkeypatch.setattr(config, "HISTORY_XLSX_PATH", str(p))
    monkeypatch.setattr(xlsx, "LOCK_TIMEOUT", 0.2)
    return p


def upload_and_reconcile(client, fixture) -> int:
    r = client.post("/api/batches", files=files_from(fixture))
    assert r.status_code == 201, r.text
    batch_id = r.json()["id"]
    assert client.post(f"/api/batches/{batch_id}/reconcile").status_code == 200
    return batch_id


def counts(db, smtp_) -> dict:
    return {"queue": db.scalar(select(func.count()).select_from(EmailQueue)),
            "items": db.scalar(select(func.count()).select_from(EmailItem)),
            "history": db.scalar(select(func.count()).select_from(EmailHistory)),
            "smtp": len(smtp_.sent)}


def db_status_counts(db) -> dict:
    return dict(db.execute(select(EmailHistory.status, func.count()).group_by(EmailHistory.status)).all())


def file_rows(path) -> list[tuple]:
    rows = list(load_workbook(path).active.iter_rows(values_only=True))
    utr, status = rows[0].index("UTR"), rows[0].index("Status")
    return [(r[utr], r[status]) for r in rows[1:]]


def result(client, batch_id, invoice_no) -> dict:
    return next(r for r in client.get(f"/api/batches/{batch_id}/results", params={"limit": 500}).json()
                if r["invoice_no"] == invoice_no)


def test_step5_upload_reconcile_queue_send_and_preview_cells(client, db, hist, smtp, monkeypatch):
    monkeypatch.setenv("DRY_RUN", "false")
    batch_id = upload_and_reconcile(client, fx)

    # reconcile alone created the queue: one email per vendor + send date, nothing else was called
    queue = client.get("/api/email-queue", params={"batch_id": batch_id}).json()
    assert sorted((q["send_on_date"], q["invoice_count"], q["status"]) for q in queue) == sorted(
        (str(fx.SEND_ON_DATES[paid]), len(invoices), "PENDING") for paid, invoices in fx.EMAIL_GROUPS.items())

    r = client.post("/api/scheduler/run-now")
    assert r.status_code == 200 and r.json()["sent"] == 2 and r.json()["failed"] == 0
    assert {m["To"] for m in smtp.sent} == {fx.VENDOR["email"]}

    expected = {str(fx.SEND_ON_DATES[paid]): [fx.EMAIL_HEADER] + [fx.EMAIL_ROWS[i] for i in invoices]
                for paid, invoices in fx.EMAIL_GROUPS.items()}
    for q in queue:
        p = client.get(f"/api/email-queue/{q['id']}/preview")
        assert p.status_code == 200
        assert table(p.json()["html"]) == expected[q["send_on_date"]]  # cell by cell, empty cells included
    sent_tables = sorted(table(m.get_body(("html",)).get_content()) for m in smtp.sent)
    assert sent_tables == sorted(expected.values())  # what went out is what the preview shows


def test_step6_rerun_reconcile_and_scheduler_add_nothing(client, db, hist, smtp, monkeypatch):
    batch_id = upload_and_reconcile(client, fx)
    for _ in range(2):  # DRY_RUN is the default: renders, logs, sends nothing, returns rows to PENDING
        assert client.post("/api/scheduler/run-now").json()["dry_run"] == 2
    assert db_status_counts(db) == {"DRY_RUN": 3} and smtp.sent == []  # one DRY_RUN row per UTR, not per run

    monkeypatch.setenv("DRY_RUN", "false")
    assert client.post("/api/scheduler/run-now").json()["sent"] == 2
    before = counts(db, smtp)
    assert client.post(f"/api/batches/{batch_id}/reconcile").status_code == 200
    for _ in range(2):
        assert client.post("/api/scheduler/run-now").json() == {"sent": 0, "failed": 0, "dry_run": 0, "not_sent": 0}
    assert counts(db, smtp) == before
    assert db_status_counts(db) == {"DRY_RUN": 3, "SENT": 3}


def test_step7_history_matches_in_db_file_and_page(client, db, hist, smtp, monkeypatch):
    upload_and_reconcile(client, fx)
    upload_and_reconcile(client, sc)  # its exceptions and manual-review rows are NOT_SENT; its emails are due later
    run_due(db, date(2025, 9, 15))  # dry run of the email due on the 15th
    monkeypatch.setenv("DRY_RUN", "false")
    smtp.fail_send = smtplib.SMTPServerDisconnected("connection lost after DATA")
    assert run_due(db, date(2025, 9, 15))["failed"] == 1
    smtp.fail_send = None
    with FileLock(f"{hist}.lock"):  # someone holds the history file: sending must not wait for it
        assert run_due(db, date(2025, 9, 17))["sent"] == 1

    in_db = db_status_counts(db)
    assert set(in_db) == {"SENT", "FAILED", "DRY_RUN", "NOT_SENT"}
    db_rows = Counter(tuple(r) for r in db.execute(select(EmailHistory.utr, EmailHistory.status)))
    assert sum(r[1] == "SENT" for r in file_rows(hist)) == 0  # the locked append was skipped, not blocked

    x = client.get("/api/email-history/export.xlsx")  # catches the file up first
    assert x.status_code == 200
    assert Counter(file_rows(hist)) == db_rows

    page = client.get("/api/email-history", params={"limit": 500})
    assert {s: n for s, n in json.loads(page.headers["x-status-counts"]).items() if n} == in_db
    assert int(page.headers["x-total-count"]) == sum(in_db.values())
    assert Counter((h["utr"], h["status"]) for h in page.json()) == db_rows

    hist.write_bytes(b"not a workbook")  # damaged: rebuild restores it from the database
    r = client.post("/api/email-history/rebuild-xlsx")
    assert r.status_code == 200 and r.json() == {"rows": sum(in_db.values())}
    assert Counter(file_rows(hist)) == db_rows


def test_step8_manual_review_and_fixes_queue_with_send_date_and_audit(client, db, hist, smtp):
    batch_id = upload_and_reconcile(client, sc)
    audits = lambda: db.scalar(select(func.count()).select_from(AuditLog))  # noqa: E731

    # no reason -> manual review, no email
    s06 = result(client, batch_id, "INV-S06")
    assert (s06["status"], s06["exception_type"], s06["send_on_date"]) == ("MANUAL_REVIEW", "REASON_MISSING", None)
    assert s06["id"] in {r["id"] for r in client.get("/api/manual-review").json()}
    client.post("/api/scheduler/run-now")
    s06_history = db.scalars(select(EmailHistory.status).where(EmailHistory.invoice_no == "INV-S06")).all()
    assert s06_history == ["NOT_SENT"]

    # reason typed on the manual review page
    before = audits()
    r = client.put(f"/api/manual-review/{s06['id']}/reason", json={"reason": "Short supply (DN-0901)"})
    assert r.status_code == 200, r.text
    assert (r.json()["status"], r.json()["reason"], r.json()["send_on_date"]) == ("DEDUCTION", "Short supply (DN-0901)", SC_SEND_ON)
    assert audits() > before

    # updated GDN file gives INV-S22 its narration
    before = audits()
    gdn = make_xlsx([{"Date": "2025-09-11 00:00:00", "Particulars": "Kiran Medicos", "Voucher No.": "722",
                      "Voucher Ref. No.": "INV-S22", "Narration": "Rate difference", "Gross Total": "1000"}])
    r = client.post("/api/manual-review/gdn-upload", files={"file": ("gdn.xlsx", gdn)})
    assert r.status_code == 200 and r.json() == {"rows_read": 1, "requeued": 1}, r.text
    s22 = result(client, batch_id, "INV-S22")
    assert (s22["status"], s22["reason"], s22["send_on_date"]) == ("DEDUCTION", "Rate difference", SC_SEND_ON)
    assert audits() > before

    # invalid vendor email fixed -> that vendor's exception rows are queued
    s07 = next(r for r in client.get("/api/exceptions").json() if r["invoice_no"] == "INV-S07")
    assert s07["exception_type"] == "INVALID_EMAIL" and s07["vendor_id"] is not None
    before = audits()
    r = client.put(f"/api/vendors/{s07['vendor_id']}/email", json={"email": "accounts@harkrishan.example"})
    assert r.status_code == 200, r.text
    s07 = result(client, batch_id, "INV-S07")
    assert (s07["status"], s07["send_on_date"]) == ("PAID_IN_FULL", SC_SEND_ON)
    assert audits() > before
