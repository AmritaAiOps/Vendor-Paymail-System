import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.contract import GdnRow
from app.models import AuditLog, ConsolidatedEntry, EmailHistory, ReconciliationResult, Vendor
from app.reports.results import HEADERS
from scripts.seed_fixture import seed


def sheet(content):
    return [[c.value for c in r] for r in load_workbook(io.BytesIO(content)).active.iter_rows()]


def result_for(db, invoice):
    return db.scalar(select(ReconciliationResult).join(ConsolidatedEntry, ConsolidatedEntry.id == ReconciliationResult.consolidated_entry_id)
                     .where(ConsolidatedEntry.invoice_no == invoice))


def set_status(db, invoice, status, exc=None):
    r = result_for(db, invoice)
    r.status, r.exception_type = status, exc
    db.commit()
    return r


@pytest.fixture
def calls(monkeypatch):
    """Stand-ins for Person A's functions: record the arguments, do nothing."""
    log = {"reason": [], "vendor": []}
    monkeypatch.setattr("app.reconcile.requeue_after_reason", lambda *a: log["reason"].append(a))
    monkeypatch.setattr("app.reconcile.requeue_vendor_exceptions", lambda *a: log["vendor"].append(a) or 0)
    return log


def test_exceptions_list_and_export(client, db):
    seed(db)
    set_status(db, "INV-2305", "EXCEPTION", "UTR_NOT_FOUND")
    set_status(db, "INV-2310", "EXCEPTION", "BOUNCED")
    r = client.get("/api/exceptions")
    assert {x["invoice_no"]: x["exception_type"] for x in r.json()} == {"INV-2305": "UTR_NOT_FOUND", "INV-2310": "BOUNCED"}
    assert "x-stub" not in r.headers and all(x["send_on_date"] is None for x in r.json())
    x = client.get("/api/exceptions/export.xlsx")
    rows = sheet(x.content)
    assert rows[0] == HEADERS and len(rows) == 3 and "x-stub" not in x.headers
    first = dict(zip(HEADERS, rows[1]))
    assert first["Invoice Amount"] == "1,20,000.00" and first["Exception Type"] == "UTR_NOT_FOUND"
    assert first["TDS Amount"] is None and first["Send Date"] is None  # empty stays empty


def test_batch_report_has_every_status(client, db):
    batch_id = seed(db)
    set_status(db, "INV-2305", "EXCEPTION", "UTR_NOT_FOUND")
    rows = sheet(client.get(f"/api/batches/{batch_id}/report.xlsx").content)
    assert rows[0] == HEADERS and sorted(r[1] for r in rows[1:]) == ["INV-2291", "INV-2305", "INV-2310"]
    assert client.get("/api/batches/999/report.xlsx").status_code == 404


def test_manual_review_list(client, db):
    seed(db)
    set_status(db, "INV-2291", "MANUAL_REVIEW", "REASON_MISSING")
    assert [x["invoice_no"] for x in client.get("/api/manual-review").json()] == ["INV-2291"]


def test_reason_is_trimmed_validated_and_requeued(client, db, calls):
    seed(db)
    r = set_status(db, "INV-2291", "MANUAL_REVIEW", "REASON_MISSING")
    ok = client.put(f"/api/manual-review/{r.id}/reason", json={"reason": "  Goods returned (DN-0457)  "})
    assert ok.status_code == 200 and "x-stub" not in ok.headers
    assert calls["reason"] == [(db, r.id, "Goods returned (DN-0457)")]
    entry = db.scalars(select(AuditLog).order_by(AuditLog.id.desc())).first()
    assert f"reconciliation_result {r.id}" in entry.what and "from " in entry.what and entry.new_value == "Goods returned (DN-0457)"

    calls["reason"].clear()
    assert client.put(f"/api/manual-review/{r.id}/reason", json={"reason": "   "}).status_code == 422
    assert client.put(f"/api/manual-review/{r.id}/reason", json={"reason": "x" * 501}).status_code == 422
    assert client.put("/api/manual-review/9999/reason", json={"reason": "x"}).status_code == 404
    paid = result_for(db, "INV-2305")
    assert client.put(f"/api/manual-review/{paid.id}/reason", json={"reason": "x"}).status_code == 409
    assert calls["reason"] == []  # none of the rejected requests reached the engine


def test_reason_refused_by_the_engine_is_a_422(client, db, monkeypatch):
    seed(db)
    r = set_status(db, "INV-2291", "MANUAL_REVIEW", "AMOUNT_MISMATCH")

    def refuse(*a):
        raise ValueError("no deduction a reason could explain")

    monkeypatch.setattr("app.reconcile.requeue_after_reason", refuse)
    res = client.put(f"/api/manual-review/{r.id}/reason", json={"reason": "x"})
    assert res.status_code == 422 and "no deduction" in res.json()["detail"]


def gdn(*rows):
    return [GdnRow(vendor_name=v, invoice_no=i, voucher_no=None, gdn_date=None, amount=Decimal("1"), reason=t) for v, i, t in rows]


def upload(client, name="gdn.xlsx"):
    return client.post("/api/manual-review/gdn-upload", files={"file": (name, b"x", "application/octet-stream")})


def test_gdn_upload_matches_review_items_and_requeues(client, db, calls, monkeypatch):
    batch_id = seed(db)
    r = set_status(db, "INV-2291", "MANUAL_REVIEW", "REASON_MISSING")
    other = ConsolidatedEntry(batch_id=batch_id, vendor_name="Other Vendor", utr="U9", invoice_no="INV-2291",
                              invoice_amount=Decimal("1"), cheque_amount=Decimal("1"))
    db.add(other)
    db.flush()
    twin = ReconciliationResult(consolidated_entry_id=other.id, invoice_value=Decimal("1"), status="MANUAL_REVIEW",
                                payment_date=date(2025, 9, 1))
    db.add(twin)
    db.commit()
    monkeypatch.setattr("app.ingest.load_gdn_file", lambda f: gdn(
        ("ABC Pharma Distributors", "inv-2291", "Goods returned (DN-0457)"),
        ("ABC PHARMA DISTRIBUTORS", "INV-2291", "Goods returned (DN-0457)"),  # same text twice is kept once
        ("ABC Pharma Distributors", "INV-2291", "Short supply"),
        ("ABC Pharma Distributors", "INV-0000", "no such invoice"),
        ("ABC Pharma Distributors", "INV-2305", "paid in full, not in review"),
        ("ABC Pharma Distributors", "INV-2291", "  "),  # blank narration ignored
    ))
    res = upload(client)
    assert res.json() == {"rows_read": 6, "requeued": 1} and "x-stub" not in res.headers
    assert calls["reason"] == [(db, r.id, "Goods returned (DN-0457); Short supply")]  # the other vendor's twin is untouched
    assert "GDN file 'gdn.xlsx'" in db.scalars(select(AuditLog).order_by(AuditLog.id.desc())).first().what


def test_gdn_upload_rejects_bad_files(client, db, calls, monkeypatch):
    assert upload(client, "gdn.csv").status_code == 422

    def not_gdn(f):
        raise ValueError("not a GDN Report")

    monkeypatch.setattr("app.ingest.load_gdn_file", not_gdn)
    res = upload(client)
    assert res.status_code == 422 and res.json()["detail"] == "not a GDN Report"
    assert calls["reason"] == []


def test_vendor_email_fix_validates_saves_and_requeues(client, db, calls):
    seed(db)
    v = db.scalar(select(Vendor))
    ok = client.put(f"/api/vendors/{v.id}/email", json={"email": " a@x.com; b@y.com, A@x.com "})
    assert ok.status_code == 200 and "x-stub" not in ok.headers
    assert ok.json()["email"] == "a@x.com, b@y.com" and ok.json()["is_email_valid"] is True
    assert calls["vendor"] == [(db, v.id)]
    entry = db.scalars(select(AuditLog).order_by(AuditLog.id.desc())).first()
    assert entry.new_value == "a@x.com, b@y.com" and entry.old_value == "accounts@abcpharma.example"

    calls["vendor"].clear()
    for bad in ("not-an-email", "a@x.com, oops", "a@x.com b@y.com"):
        r = client.put(f"/api/vendors/{v.id}/email", json={"email": bad})
        assert r.status_code == 422, bad
    assert client.put("/api/vendors/9999/email", json={"email": "a@x.com"}).status_code == 404
    db.refresh(v)
    assert v.email == "a@x.com, b@y.com" and calls["vendor"] == []


def test_vendor_email_is_kept_if_requeue_fails(client, db, monkeypatch):
    seed(db)
    v = db.scalar(select(Vendor))

    def boom(*a):
        raise RuntimeError("engine down")

    monkeypatch.setattr("app.reconcile.requeue_vendor_exceptions", boom)
    r = client.put(f"/api/vendors/{v.id}/email", json={"email": "new@x.com"})
    assert r.status_code == 500 and "was saved" in r.json()["detail"]
    db.refresh(v)
    assert v.email == "new@x.com"


def test_filtered_history_export(client, db, tmp_path, monkeypatch):
    monkeypatch.setattr("app.config.HISTORY_XLSX_PATH", str(tmp_path / "h.xlsx"))
    for vendor, status, day in [("ABC Pharma", "SENT", 12), ("Zed Labs", "NOT_SENT", 14), ("Zed Labs", "FAILED", 15)]:
        db.add(EmailHistory(vendor_name=vendor, status=status, utr=f"U{day}", logged_at=datetime(2025, 9, day, 10, 0)))
    db.commit()
    rows = sheet(client.get("/api/email-history/export.xlsx", params={"vendor": "zed", "status": "FAILED"}).content)
    assert [r[6] for r in rows[1:]] == ["U15"]
    rows = sheet(client.get("/api/email-history/export.xlsx", params={"date_from": "2025-09-13"}).content)
    assert [r[6] for r in rows[1:]] == ["U15", "U14"]  # newest first
    assert len(sheet(client.get("/api/email-history/export.xlsx").content)) == 4  # no filter: the whole file
