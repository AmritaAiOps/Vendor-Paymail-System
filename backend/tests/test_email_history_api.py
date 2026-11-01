import io
import json
from datetime import datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import config
from app.email_builder import build_queue
from app.models import EmailHistory, EmailQueue
from scripts.seed_fixture import seed


@pytest.fixture(autouse=True)
def history_file(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "HISTORY_XLSX_PATH", str(tmp_path / "h.xlsx"))


def add(db, vendor, status, day, utr, invoice="INV-1", **kw):
    h = EmailHistory(vendor_name=vendor, status=status, utr=utr, invoice_no=invoice, logged_at=datetime(2025, 9, day, 10, 0), **kw)
    db.add(h)
    db.commit()
    return h.id


def test_history_filters_pagination_and_counts(client, db):
    add(db, "ABC Pharma", "SENT", 12, "UTR-A1", invoice_amount=Decimal("545060"))
    add(db, "ABC Pharma", "DRY_RUN", 13, "UTR-A2")
    add(db, "Zed Labs", "NOT_SENT", 14, "UTR-Z1", invoice="INV-99", not_sent_reason="UTR_NOT_FOUND")
    add(db, "Zed Labs", "FAILED", 15, "UTR-Z2", error_message="boom")
    r = client.get("/api/email-history")
    assert [h["utr"] for h in r.json()] == ["UTR-Z2", "UTR-Z1", "UTR-A2", "UTR-A1"]  # newest first
    assert r.headers["x-total-count"] == "4" and "x-stub" not in r.headers
    assert json.loads(r.headers["x-status-counts"]) == {"SENT": 1, "FAILED": 1, "DRY_RUN": 1, "NOT_SENT": 1}

    p = client.get("/api/email-history", params={"limit": 2, "offset": 2})
    assert [h["utr"] for h in p.json()] == ["UTR-A2", "UTR-A1"] and p.headers["x-total-count"] == "4"

    s = client.get("/api/email-history", params={"status": "FAILED"})
    assert [h["utr"] for h in s.json()] == ["UTR-Z2"] and s.headers["x-total-count"] == "1"
    assert json.loads(s.headers["x-status-counts"])["SENT"] == 1  # counts ignore the status filter

    d = client.get("/api/email-history", params={"date_from": "2025-09-13", "date_to": "2025-09-14"})
    assert [h["utr"] for h in d.json()] == ["UTR-Z1", "UTR-A2"]
    assert [h["utr"] for h in client.get("/api/email-history", params={"vendor": "zed"}).json()] == ["UTR-Z2", "UTR-Z1"]
    assert [h["utr"] for h in client.get("/api/email-history", params={"q": "inv-99"}).json()] == ["UTR-Z1"]
    assert [h["utr"] for h in client.get("/api/email-history", params={"q": "utr-a"}).json()] == ["UTR-A2", "UTR-A1"]
    assert client.get("/api/email-history", params={"q": "%"}).json() == []  # wildcard is escaped, not a match-all
    assert client.get("/api/email-history", params={"status": "BOGUS"}).status_code == 422


def test_history_row_and_export(client, db):
    assert client.get("/api/email-history/export.xlsx").status_code == 404  # nothing logged yet
    hid = add(db, "ABC Pharma", "SENT", 12, "UTR-A1", invoice_amount=Decimal("545060"))
    row = client.get(f"/api/email-history/{hid}")
    assert row.status_code == 200 and row.json()["invoice_amount"] == "545060.00" and "x-stub" not in row.headers
    assert client.get("/api/email-history/999").status_code == 404

    x = client.get("/api/email-history/export.xlsx")
    assert x.status_code == 200 and "x-stub" not in x.headers
    sheet = [[c.value for c in r] for r in load_workbook(io.BytesIO(x.content)).active.iter_rows()]
    assert len(sheet) == 2 and sheet[1][5] == "5,45,060.00"


def test_preview_shows_the_email_as_sent(client, db):
    seed(db)
    from app.models import ReconciliationResult, Vendor
    db.execute(ReconciliationResult.__table__.update().values(vendor_id=db.scalar(select(Vendor.id))))
    db.commit()
    build_queue(db)
    qid = db.scalar(select(EmailQueue.id))
    live = client.get(f"/api/email-queue/{qid}/preview").json()
    q = db.get(EmailQueue, qid)
    q.status, q.subject, q.body_html = "SENT", "As sent", "<p>as sent</p>"
    db.commit()
    sent = client.get(f"/api/email-queue/{qid}/preview").json()
    assert sent["html"] == "<p>as sent</p>" and sent["subject"] == "As sent" and live["html"] != sent["html"]
