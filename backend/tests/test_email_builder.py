import io
from decimal import Decimal
from html.parser import HTMLParser

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.email_builder import build_queue, render_email
from app.email_builder.rows import inr
from app.models import EmailItem, EmailQueue, ReconciliationResult
from scripts.seed_fixture import seed
from tests.fixtures import three_row_example as fx


class _Table(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self._cell = [], None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.rows.append([])
        elif tag in ("td", "th"):
            self._cell = ""

    def handle_endtag(self, tag):
        if tag in ("td", "th"):
            self.rows[-1].append(self._cell)
            self._cell = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell += data


def table(html):
    p = _Table()
    p.feed(html)
    return p.rows


def test_inr_grouping():
    assert inr(Decimal("545060")) == "5,45,060.00"
    assert inr(Decimal("999.5")) == "999.50"
    assert inr(Decimal("12345678.9")) == "1,23,45,678.90"
    assert inr(None) == ""


def test_same_vendor_same_day_is_one_email(db):
    seed(db)
    ids = build_queue(db)
    counts = sorted(db.scalar(select(func.count()).where(EmailItem.email_queue_id == i)) for i in ids)
    assert counts == [1, 2]  # INV-2305 + INV-2310 together, INV-2291 alone
    assert {q.send_on_date for q in db.scalars(select(EmailQueue))} == set(fx.SEND_ON_DATES.values())


def test_build_queue_twice_changes_nothing(db):
    batch_id = seed(db)
    build_queue(db, batch_id=batch_id)
    before = (db.scalar(select(func.count(EmailQueue.id))), db.scalar(select(func.count(EmailItem.id))))
    assert build_queue(db, batch_id=batch_id) == []
    all_ids = list(db.scalars(select(ReconciliationResult.id)))
    assert build_queue(db, result_ids=all_ids) == []
    assert (db.scalar(select(func.count(EmailQueue.id))), db.scalar(select(func.count(EmailItem.id)))) == before


def test_exception_results_are_not_queued(db):
    seed(db)
    for r in db.scalars(select(ReconciliationResult)):
        r.status = "MANUAL_REVIEW"
    db.commit()
    assert build_queue(db) == []


def test_rendered_cells_match_fixture(db):
    seed(db)
    build_queue(db)
    for q in db.scalars(select(EmailQueue).order_by(EmailQueue.send_on_date)):
        subject, html, text = render_email(db, q.id)
        rows = table(html)
        assert subject == "Payment Details – ABC Pharma Distributors – Sudhamayi Enterprises Pvt. Ltd."
        assert rows[0] == fx.EMAIL_HEADER
        expected = [fx.EMAIL_ROWS[i] for i in fx.EMAIL_GROUPS[next(d for d, s in fx.SEND_ON_DATES.items() if s == q.send_on_date)]]
        assert rows[1:] == expected  # includes the empty cells
        assert text.startswith("Dear ABC Pharma Distributors,") and all(r[0] in text for r in expected)
        assert q.subject == subject and q.body_html == html


def test_vendor_text_is_escaped(db):
    seed(db)
    db.execute(ReconciliationResult.__table__.update().where(ReconciliationResult.invoice_deduction.is_not(None))
               .values(reason="<script>alert(1)</script>"))
    db.commit()
    build_queue(db)
    html = "".join(render_email(db, q)[1] for q in db.scalars(select(EmailQueue.id)))
    assert "<script>" not in html and "&lt;script&gt;" in html


def test_queue_preview_and_proforma_endpoints(client, db):
    batch_id = seed(db)
    build_queue(db)
    r = client.get("/api/email-queue", params={"batch_id": batch_id})
    assert "x-stub" not in r.headers and sorted(q["invoice_count"] for q in r.json()) == [1, 2]
    assert client.get("/api/email-queue", params={"status": "SENT"}).json() == []
    p = client.get(f"/api/email-queue/{r.json()[0]['id']}/preview")
    assert "x-stub" not in p.headers and "<table" in p.json()["html"]
    assert client.get("/api/email-queue/999/preview").status_code == 404

    x = client.get(f"/api/batches/{batch_id}/proforma.xlsx")
    assert "x-stub" not in x.headers
    cells = [[c.value or "" for c in row] for row in load_workbook(io.BytesIO(x.content)).active.iter_rows()]
    assert ["Vendor Name", "ABC Pharma Distributors"] == cells[0][:2]
    assert fx.EMAIL_HEADER in cells and all(v in cells for v in fx.EMAIL_ROWS.values())
    assert client.get("/api/batches/999/proforma.xlsx").status_code == 404
