"""The one xlsx helper behind every report and export: a styled sheet, rows appended as safe cells, a download response."""
import io

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook(title: str, headers: list[str], widths: list[int] | None = None) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return wb


def append(ws, cells: list) -> None:
    """None stays an empty cell. Text is forced to text so a value starting with "=" can't become a formula."""
    ws.append(cells)
    for c in ws[ws.max_row]:
        if isinstance(c.value, str):
            c.data_type = "s"


def response(wb: Workbook, filename: str) -> Response:
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{filename}"'})
