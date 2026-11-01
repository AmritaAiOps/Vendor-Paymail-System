"""Append-only Excel copy of email_history at config.HISTORY_XLSX_PATH. Owner: Person B.

The database is the source of truth. The "last email_history id written" mark lives in the workbook's own
properties (description), saved in the same file write as the rows: rows and mark succeed or fail together, so a
crash can't leave rows written with the mark behind (a double append). Delete the file and the next call refills it.
"""
import logging
import os
import re
from pathlib import Path

from filelock import FileLock  # shortcut: filelock package, the stdlib has no cross-platform file lock; see requirements.txt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import config
from app.email_builder.rows import day, inr
from app.models import EmailHistory

log = logging.getLogger(__name__)

COLUMNS = ["Logged At", "Vendor", "To", "Subject", "Invoice No", "Invoice Amount", "UTR", "Payment Date",
           "Amount Paid", "TDS Amount", "Invoice Deduction", "Reason", "Status", "Not Sent Reason",
           "Scheduled Send Date", "Sent At", "Error"]
LOCK_TIMEOUT = 10  # seconds to wait for another writer before giving up until the next call
_MARK = re.compile(r"history_last_id=(\d+)")


def _stamp(dt) -> str:
    return dt.strftime("%d-%b-%y %H:%M:%S") if dt else ""


def _cells(h: EmailHistory) -> list:
    """Cells in COLUMNS order. Empty means a truly empty cell (None), never "0" or "n/a"."""
    cells = [_stamp(h.logged_at), h.vendor_name, h.to_email, h.subject, h.invoice_no, inr(h.invoice_amount), h.utr,
             day(h.payment_date), inr(h.amount_paid), inr(h.tds_amount) if h.tds_amount else "",
             inr(h.invoice_deduction) if h.invoice_deduction else "", h.reason, str(h.status),
             str(h.not_sent_reason) if h.not_sent_reason else "", day(h.scheduled_send_date), _stamp(h.sent_at),
             h.error_message]
    return [c or None for c in cells]


def _new_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Email History"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJKLMNOPQ", (20, 28, 30, 44, 14, 16, 22, 14, 16, 14, 18, 28, 12, 18, 18, 20, 36)):
        ws.column_dimensions[col].width = w
    return wb


def _last_id(wb: Workbook) -> int:
    m = _MARK.fullmatch(wb.properties.description or "")
    if not m:  # a file we can't place: appending everything could duplicate rows, so make a human rebuild it
        raise ValueError("history file has no id mark; run POST /api/email-history/rebuild-xlsx")
    return int(m[1])


def _write(session: Session, rebuild: bool) -> int:
    path = Path(config.HISTORY_XLSX_PATH)
    path.parent.mkdir(parents=True, exist_ok=True)
    with FileLock(f"{path}.lock", timeout=LOCK_TIMEOUT):
        fresh = rebuild or not path.exists()
        wb = _new_workbook() if fresh else load_workbook(path)
        last = 0 if fresh else _last_id(wb)
        # shortcut: a row committed after a higher id (concurrent writers) is missed until the next rebuild
        rows = session.scalars(select(EmailHistory).where(EmailHistory.id > last).order_by(EmailHistory.id)).all()
        if not rows and not rebuild:
            return 0
        ws = wb.active
        for h in rows:
            ws.append(_cells(h))
            for c in ws[ws.max_row]:
                if isinstance(c.value, str):
                    c.data_type = "s"  # vendor-supplied text starting "=" must not become a formula
        wb.properties.description = f"history_last_id={rows[-1].id if rows else last}"
        tmp = path.with_name(path.name + ".tmp")
        try:
            wb.save(tmp)
            os.replace(tmp, path)  # all or nothing: a file open in Excel makes this fail and leaves the old file intact
        finally:
            tmp.unlink(missing_ok=True)
        return len(rows)


def append_new_rows(session: Session) -> int:
    """Append email_history rows not yet in the Excel file. Returns how many were appended.
    Never raises: a locked or broken file is logged and the next call retries from the same mark."""
    try:
        return _write(session, rebuild=False)
    except Exception:
        log.exception("email history Excel append failed; the database history is unaffected")
        return 0


def rebuild(session: Session) -> int:
    """Rewrite the Excel file from the whole email_history table. Returns the row count.
    Raises if the file can't be written, so the caller can tell the user."""
    return _write(session, rebuild=True)
