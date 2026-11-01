"""Excel proforma: per vendor a header block, then the same 8-column rows as the email."""
import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.email_builder.rows import HEADER, row, sendable

NAVY = PatternFill("solid", fgColor="1F2A44")
EDGE = Side(style="thin", color="D9DEE7")
RIGHT = {1, 4, 5, 6}  # zero-based: Invoice Amount, Amount Paid, TDS Amount, Invoice Deduction


def proforma_xlsx(session: Session, batch_id: int) -> bytes:
    by_vendor = defaultdict(list)
    for vendor, res, entry in sendable(session, batch_id=batch_id):
        by_vendor[vendor].append(row(res, entry))
    wb = Workbook()
    ws = wb.active
    ws.title = "Proforma"

    def put(r, c, v, **style):
        cell = ws.cell(r, c, v)
        cell.data_type = "s" if isinstance(v, str) else cell.data_type  # vendor text starting "=" must not become a formula
        for k, s in style.items():
            setattr(cell, k, s)

    r = 1
    for vendor, rows in by_vendor.items():
        put(r, 1, "Vendor Name", font=Font(bold=True)); put(r, 2, vendor.party_name); r += 1
        put(r, 1, "Email id", font=Font(bold=True)); put(r, 2, vendor.email or ""); r += 1
        for c, h in enumerate(HEADER, 1):
            put(r, c, h, font=Font(bold=True, color="FFFFFF"), fill=NAVY, border=Border(EDGE, EDGE, EDGE, EDGE))
        r += 1
        for i, cells in enumerate(rows):
            for c, v in enumerate(cells, 1):
                put(r, c, v, border=Border(EDGE, EDGE, EDGE, EDGE),
                    alignment=Alignment(horizontal="right" if c - 1 in RIGHT else "left"),
                    fill=PatternFill("solid", fgColor="F5F7FA") if i % 2 else PatternFill())
            r += 1
        r += 1  # blank line between vendors
    for col, w in zip("ABCDEFGH", (16, 16, 22, 14, 16, 14, 18, 32)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
